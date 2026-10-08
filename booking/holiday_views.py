import io
import secrets

from django.contrib import messages
from django.contrib.sessions.models import Session
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from booking.holiday_forms import HolidayConfirmForm, HolidayUploadForm
from booking.services.holiday_import import (
    HolidayImportChanged,
    apply_holiday_import,
    build_holiday_preview,
    parse_holiday_workbook,
)
from booking.staff_views import _locked_verified_staff, staff_required

PREVIEW_SESSION_KEY = "holiday_import_preview"
PREVIEW_SECONDS = 15 * 60


def _pending_preview(request, state):
    if not isinstance(state, dict):
        return None
    try:
        age = timezone.now().timestamp() - float(state["created_at"])
    except (KeyError, TypeError, ValueError):
        return None
    if (
        not 0 <= age < PREVIEW_SECONDS
        or state.get("actor_id") != request.user.pk
        or state.get("auth_version") != request.user.auth_version
        or not isinstance(state.get("rows"), list)
        or not state.get("preview_id")
        or not state.get("snapshot")
    ):
        return None
    return state


def _save_preview(request, rows, preview, *, filename=""):
    state = {
        "preview_id": secrets.token_urlsafe(24),
        "actor_id": request.user.pk,
        "auth_version": request.user.auth_version,
        "created_at": timezone.now().timestamp(),
        "rows": rows,
        "snapshot": preview["snapshot"],
        "filename": filename,
    }
    request.session[PREVIEW_SESSION_KEY] = state
    return state


def _render_import(
    request, *, state=None, preview=None, upload_form=None, confirm_form=None, errors=None, status=200
):
    if state is not None:
        preview = preview if preview is not None else build_holiday_preview(state["rows"])
        confirm_form = confirm_form or HolidayConfirmForm(
            initial={"preview_id": state["preview_id"]},
            affected_bookings=preview["summary"]["affected_bookings"],
        )
    return render(
        request,
        "booking/staff_holiday_import.html",
        {
            "upload_form": upload_form if upload_form is not None else HolidayUploadForm(),
            "preview": preview,
            "preview_filename": state.get("filename", "") if state else "",
            "confirm_form": confirm_form,
            "errors": errors or [],
        },
        status=status,
    )


@staff_required
@require_http_methods(["GET", "POST"])
def holiday_import(request):
    if request.method == "GET":
        state = _pending_preview(request, request.session.get(PREVIEW_SESSION_KEY))
        if state is None:
            request.session.pop(PREVIEW_SESSION_KEY, None)
        else:
            preview = build_holiday_preview(state["rows"])
            if preview["snapshot"] != state["snapshot"]:
                state = _save_preview(request, state["rows"], preview, filename=state.get("filename", ""))
                messages.warning(
                    request, "The holiday list or existing meetings changed. Review the refreshed preview."
                )
            return _render_import(request, state=state, preview=preview)
        return _render_import(request)

    action = request.POST.get("action", "")
    if action in {"discard", "cancel"}:
        request.session.pop(PREVIEW_SESSION_KEY, None)
        return redirect("staff-holiday-import" if action == "discard" else "staff-holidays")
    if action == "preview":
        request.session.pop(PREVIEW_SESSION_KEY, None)
        form = HolidayUploadForm(request.POST, request.FILES)
        if not form.is_valid():
            return _render_import(request, upload_form=form)
        result = parse_holiday_workbook(form.cleaned_data["file"])
        if result["errors"]:
            return _render_import(request, errors=result["errors"])
        preview = build_holiday_preview(result["rows"])
        state = _save_preview(request, result["rows"], preview, filename=form.cleaned_data["file"].name)
        return _render_import(request, state=state, preview=preview)

    if action != "confirm":
        return _render_import(
            request,
            errors=[{"row_number": None, "message": "Choose a file to preview before importing."}],
            status=400,
        )

    with transaction.atomic():
        actor = _locked_verified_staff(request)
        if actor is None:
            return redirect("staff-login")
        # Read the persisted preview under a session lock. Two concurrent confirmations
        # cannot reuse the same preview after the first request consumes it.
        session_record = (
            Session.objects.select_for_update().filter(session_key=request.session.session_key).first()
        )
        stored = session_record.get_decoded() if session_record else {}
        state = _pending_preview(request, stored.get(PREVIEW_SESSION_KEY))
        if (
            state is None
            or stored.get("staff_verified") is not True
            or stored.get("staff_auth_version") != actor.auth_version
            or str(stored.get("_auth_user_id")) != str(actor.pk)
        ):
            request.session.pop(PREVIEW_SESSION_KEY, None)
            return _render_import(
                request,
                errors=[
                    {
                        "row_number": None,
                        "message": "This preview expired or has already been used. Upload the file again.",
                    }
                ],
                status=400,
            )

        submitted_id = request.POST.get("preview_id", "")
        if len(submitted_id) > 64 or not secrets.compare_digest(
            submitted_id.encode(), str(state["preview_id"]).encode()
        ):
            # A stale tab must not invalidate a newer preview in the same session.
            request.session[PREVIEW_SESSION_KEY] = state
            return _render_import(
                request,
                errors=[
                    {
                        "row_number": None,
                        "message": "This confirmation belongs to an older preview. Open the import page to review the current file.",
                    }
                ],
                status=400,
            )

        preview = build_holiday_preview(state["rows"])
        if preview["snapshot"] != state["snapshot"]:
            state = _save_preview(request, state["rows"], preview, filename=state.get("filename", ""))
            request.session.save()
            messages.warning(
                request,
                "The holiday list or existing meetings changed. Review the refreshed preview and confirm again.",
            )
            return _render_import(request, state=state, preview=preview)

        form = HolidayConfirmForm(request.POST, affected_bookings=preview["summary"]["affected_bookings"])
        if not form.is_valid():
            return _render_import(request, state=state, preview=preview, confirm_form=form)
        try:
            counts = apply_holiday_import(state["rows"], actor, state["snapshot"])
        except HolidayImportChanged:
            preview = build_holiday_preview(state["rows"])
            state = _save_preview(request, state["rows"], preview, filename=state.get("filename", ""))
            request.session.save()
            messages.warning(
                request,
                "The holiday list or existing meetings changed. Review the refreshed preview and confirm again.",
            )
            return _render_import(request, state=state, preview=preview)
        request.session.pop(PREVIEW_SESSION_KEY, None)
        request.session.save()
        messages.success(
            request,
            f"Holidays imported: {counts['new']} added, {counts['update']} updated, {counts['unchanged']} unchanged. Existing meetings are kept.",
        )
    return redirect("staff-holidays")


@staff_required
@require_GET
def holiday_template(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Holidays"
    sheet.append(["Date", "Occasion"])
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 60
    sheet.freeze_panes = "A2"
    sheet["A2"].number_format = "yyyy-mm-dd"
    sheet["A1"].comment = Comment(
        "Use a Gregorian (AD) Excel date or YYYY-MM-DD text. One date per row.", "MBS"
    )
    sheet["B1"].comment = Comment("Enter the holiday name, up to 160 characters.", "MBS")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163F64")
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="MBS_Holiday_Template.xlsx"'
    return response
