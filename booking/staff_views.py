import io
import logging
import re
from datetime import date, datetime, timedelta
from functools import wraps
from zoneinfo import ZoneInfo

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.password_validation import validate_password
from django.contrib.auth.tokens import default_token_generator
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.encoding import force_str
from django.utils.http import urlsafe_base64_decode, urlsafe_base64_encode
from django.views.decorators.cache import never_cache
from django.views.decorators.debug import sensitive_post_parameters, sensitive_variables
from django.views.decorators.http import require_GET, require_http_methods, require_POST
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from booking.forms import BookingRejectionForm, PolicyForm, RoomBlockForm, RoomForm
from booking.management_forms import UserAccessForm
from booking.models import (
    AuditEvent,
    BookingPolicy,
    CompanyHoliday,
    EmailToken,
    Notification,
    Reservation,
    Room,
    RoomFacility,
    User,
)
from booking.services.auth_security import staff_session_verified
from booking.services.bookings import BookingError, approve_booking, reject_booking
from booking.services.checkin import CheckInError, manual_checkin
from booking.services.mail_delivery import deliver_mail
from booking.services.reporting import report_data
from booking.services.reservations import ReservationConflict, reserve_time
from booking.services.room_photos import delete_room_photo
from booking.services.staff_auth import finish_staff_login, start_staff_login
from booking.ui import paginate

logger = logging.getLogger(__name__)


def staff_required(view):
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if not staff_session_verified(request):
            return redirect("staff-login")
        return view(request, *args, **kwargs)

    return never_cache(wrapped)


@never_cache
@sensitive_post_parameters("password")
@require_http_methods(["GET", "POST"])
def staff_login(request):
    if request.method == "POST":
        if start_staff_login(request.POST.get("email"), request.POST.get("password"), request):
            return redirect("staff-code")
        messages.error(request, "Sign-in could not be completed. Check your details or contact IT.")
    return render(
        request,
        "booking/staff_login.html",
        {"email_unavailable": settings.EMAIL_BACKEND == "booking.mail_backends.DisabledEmailBackend"},
    )


@never_cache
@sensitive_post_parameters("code")
@require_http_methods(["GET", "POST"])
def staff_code(request):
    if not request.session.get("pending_staff_token_id"):
        return redirect("staff-login")
    if request.method == "POST":
        if finish_staff_login(request.POST.get("code", ""), request):
            return redirect("staff-dashboard")
        messages.error(request, "That code is invalid or expired. Request a new code if needed.")
    return render(
        request,
        "booking/staff_code.html",
        {
            "local_mail": settings.MAIL_MODE == "file"
            and not settings.PRODUCTION_ENABLED
            and not settings.HTTPS_ENABLED
        },
    )


@staff_required
@require_GET
def dashboard(request):
    today = timezone.localdate()
    return render(
        request,
        "booking/staff_dashboard.html",
        {
            "today_bookings": Reservation.objects.filter(
                kind=Reservation.Kind.BOOKING,
                starts_at__date=today,
                status__in=["approved", "checked_in", "completed"],
            ).count(),
            "rooms_count": Room.objects.filter(is_active=True).count(),
            "mail_failed": Notification.objects.filter(status=Notification.Status.FAILED).count(),
            "pending_count": Reservation.objects.filter(
                kind=Reservation.Kind.BOOKING, status=Reservation.Status.PENDING
            ).count(),
            "no_shows": Reservation.objects.filter(
                status=Reservation.Status.NO_SHOW, starts_at__date=today
            ).count(),
            "upcoming": Reservation.objects.filter(
                kind=Reservation.Kind.BOOKING, starts_at__gte=timezone.now(), status="approved"
            )
            .select_related("room", "organizer")
            .order_by("starts_at")[:12],
        },
    )


@staff_required
@require_GET
def rooms(request):
    context = paginate(
        request,
        Room.objects.all().prefetch_related("facilities").order_by("location", "floor", "name", "pk"),
    )
    context["rooms"] = context["page_obj"]
    return render(request, "booking/staff_rooms.html", context)


@staff_required
@require_http_methods(["GET", "POST"])
@transaction.atomic
def room_form(request, room_id=None):
    queryset = Room.objects.select_for_update() if request.method == "POST" else Room.objects.all()
    room = get_object_or_404(queryset, pk=room_id) if room_id else None
    old_photo_name = room.photo.name if room else ""
    form = RoomForm(
        request.POST if request.method == "POST" else None,
        request.FILES if request.method == "POST" else None,
        instance=room,
        initial={"facilities_text": "\n".join(room.facilities.values_list("name", flat=True))}
        if room and request.method == "GET"
        else None,
    )
    if request.method == "POST" and form.is_valid():
        try:
            with transaction.atomic():
                saved = form.save()
                facilities = [
                    item.strip()
                    for item in re.split(r"[\n,;]+", form.cleaned_data["facilities_text"])
                    if item.strip()
                ]
                RoomFacility.objects.filter(room=saved).delete()
                RoomFacility.objects.bulk_create(
                    [RoomFacility(room=saved, name=name) for name in dict.fromkeys(facilities)]
                )
                AuditEvent.objects.create(
                    actor=request.user,
                    action="room_modified" if room else "room_created",
                    target_type="room",
                    target_id=saved.pk,
                    outcome="success",
                    details={"active": saved.is_active},
                )
                if old_photo_name and old_photo_name != saved.photo.name:
                    transaction.on_commit(
                        lambda: delete_room_photo(saved.photo.storage, old_photo_name), robust=True
                    )
        except IntegrityError:
            if form.instance.photo.name != old_photo_name:
                delete_room_photo(form.instance.photo.storage, form.instance.photo.name)
            form.add_error(
                None,
                "A room with this name, location and floor already exists. Please check the room directory.",
            )
        except OSError as exc:
            if form.instance.photo.name != old_photo_name:
                delete_room_photo(form.instance.photo.storage, form.instance.photo.name)
            logger.warning("Room photo storage unavailable (%s)", type(exc).__name__)
            form.add_error("photo", "The photo could not be saved. Please try again or contact IT.")
        except Exception:
            if form.instance.photo.name != old_photo_name:
                delete_room_photo(form.instance.photo.storage, form.instance.photo.name)
            raise
        else:
            messages.success(request, "Room saved")
            return redirect("staff-rooms")
    return render(
        request,
        "booking/staff_form.html",
        {
            "form": form,
            "heading": "Edit room" if room else "Add room",
            "room": room,
            "has_room_photo": bool(old_photo_name),
        },
    )


@staff_required
@require_GET
def bookings(request, pending_only=False):
    qs = (
        Reservation.objects.filter(kind=Reservation.Kind.BOOKING)
        .select_related("room", "organizer")
        .order_by("-starts_at", "-pk")
    )
    search = request.GET.get("q", "").strip()[:100]
    status = Reservation.Status.PENDING if pending_only else request.GET.get("status", "")
    if search:
        qs = qs.filter(
            Q(title__icontains=search)
            | Q(organizer__email__icontains=search)
            | Q(room__name__icontains=search)
        )
    if status in Reservation.Status.values:
        qs = qs.filter(status=status)
    if pending_only:
        qs = qs.order_by("starts_at", "created_at", "pk")
    room_id = request.GET.get("room", "")[:10]
    if room_id.isascii() and room_id.isdigit():
        qs = qs.filter(room_id=int(room_id))
    for parameter, lookup in (("from", "starts_at__date__gte"), ("to", "starts_at__date__lte")):
        value = request.GET.get(parameter, "")
        if value:
            try:
                qs = qs.filter(**{lookup: date.fromisoformat(value)})
            except ValueError:
                messages.error(request, "Enter dates in year-month-day format.")
    context = paginate(request, qs)
    context.update(
        {
            "bookings": context["page_obj"],
            "search": search,
            "status": status,
            "statuses": [
                choice
                for choice in Reservation.Status.choices
                if choice[0] not in {"blocked", "block_cancelled"}
            ],
            "pending_only": pending_only,
            "rooms": Room.objects.all(),
            "selected_room": room_id,
        }
    )
    return render(request, "booking/staff_bookings.html", context)


@staff_required
@require_POST
def booking_approve(request, booking_id):
    get_object_or_404(Reservation, pk=booking_id, kind=Reservation.Kind.BOOKING)
    try:
        approve_booking(booking_id, actor=request.user, request=request)
    except BookingError as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, "Booking approved. Confirmation has been queued for the requester.")
    return redirect("booking-detail", booking_id=booking_id)


@staff_required
@require_POST
def booking_reject(request, booking_id):
    booking = get_object_or_404(
        Reservation.objects.select_related("room", "organizer"),
        pk=booking_id,
        kind=Reservation.Kind.BOOKING,
    )
    form = BookingRejectionForm(request.POST)
    if form.is_valid():
        try:
            reject_booking(
                booking_id, actor=request.user, request=request, reason=form.cleaned_data["reason"]
            )
        except BookingError as exc:
            form.add_error(None, str(exc))
        else:
            messages.success(request, "Booking rejected. The requester can see your reason.")
            return redirect("booking-detail", booking_id=booking_id)
    return render(
        request,
        "booking/staff_form.html",
        {
            "form": form,
            "heading": f"Reject request: {booking.title}",
            "submit_label": "Reject request",
            "cancel_url": reverse("booking-detail", args=[booking.pk]),
        },
        status=400,
    )


@staff_required
@require_POST
def staff_checkin(request, booking_id):
    try:
        manual_checkin(booking_id, actor=request.user)
        messages.success(request, "Meeting checked in")
    except CheckInError as exc:
        messages.error(request, str(exc))
    return redirect("booking-detail", booking_id=booking_id)


@staff_required
@require_http_methods(["GET", "POST"])
@transaction.atomic
def block_new(request):
    form = RoomBlockForm(request.POST if request.method == "POST" else None)
    if request.method == "POST" and form.is_valid():
        data = form.cleaned_data
        start = datetime.combine(data["date"], data["start_time"], ZoneInfo("Asia/Kathmandu"))
        end = datetime.combine(data["date"], data["end_time"], ZoneInfo("Asia/Kathmandu"))
        if end <= start:
            form.add_error("end_time", "End must be after start")
        elif start <= timezone.now():
            form.add_error("date", "Choose a future period")
        else:
            try:
                block = reserve_time(
                    room_id=data["room"].pk,
                    occupied_from=start,
                    occupied_until=end,
                    starts_at=start,
                    ends_at=end,
                    kind=Reservation.Kind.BLOCK,
                    status=Reservation.Status.BLOCKED,
                    created_by=request.user,
                    block_reason=data["block_reason"],
                )
            except ReservationConflict:
                form.add_error(None, "This room already has a booking or closure during that time")
            else:
                AuditEvent.objects.create(
                    actor=request.user,
                    action="room_blocked",
                    target_type="reservation",
                    target_id=block.pk,
                    outcome="success",
                )
                messages.success(request, "Room blocked")
                return redirect("staff-blocks")
    return render(request, "booking/staff_form.html", {"form": form, "heading": "Block a room"})


@staff_required
@require_GET
def blocks(request):
    context = paginate(
        request,
        Reservation.objects.filter(kind=Reservation.Kind.BLOCK)
        .select_related("room")
        .order_by("-starts_at", "-pk"),
    )
    context["blocks"] = context["page_obj"]
    return render(request, "booking/staff_blocks.html", context)


@staff_required
@require_POST
def block_cancel(request, block_id):
    with transaction.atomic():
        block = get_object_or_404(
            Reservation.objects.select_for_update(),
            pk=block_id,
            kind=Reservation.Kind.BLOCK,
            status=Reservation.Status.BLOCKED,
        )
        block.status = Reservation.Status.BLOCK_CANCELLED
        block.save(update_fields=["status", "updated_at"])
        AuditEvent.objects.create(
            actor=request.user,
            action="room_block_cancelled",
            target_type="reservation",
            target_id=block.pk,
            outcome="success",
        )
    messages.success(request, "Room closure removed")
    return redirect("staff-blocks")


@staff_required
@require_http_methods(["GET", "POST"])
@transaction.atomic
def policy(request):
    qs = (
        BookingPolicy.objects.select_for_update() if request.method == "POST" else BookingPolicy.objects.all()
    )
    policy = qs.get(pk=1)
    form = PolicyForm(request.POST if request.method == "POST" else None, instance=policy)
    if request.method == "POST" and form.is_valid():
        saved = form.save(commit=False)
        saved.updated_by = request.user
        saved.save()
        AuditEvent.objects.create(
            actor=request.user,
            action="policy_modified",
            target_type="booking_policy",
            target_id=1,
            outcome="success",
        )
        messages.success(request, "Booking rules saved")
        return redirect("staff-policy")
    return render(request, "booking/staff_form.html", {"form": form, "heading": "Booking rules"})


@staff_required
@require_http_methods(["GET", "POST"])
@transaction.atomic
def holidays(request):
    if request.method == "POST":
        try:
            day = date.fromisoformat(request.POST.get("date", ""))
            name = request.POST.get("name", "").strip()[:160]
            if not name:
                raise ValueError
            item, created = CompanyHoliday.objects.update_or_create(date=day, defaults={"name": name})
            AuditEvent.objects.create(
                actor=request.user,
                action="holiday_saved",
                target_type="holiday",
                target_id=item.pk,
                outcome="success",
            )
            messages.success(request, "Holiday saved")
        except ValueError:
            messages.error(request, "Enter a valid date and holiday name")
        return redirect("staff-holidays")
    context = paginate(request, CompanyHoliday.objects.order_by("-date", "-pk"))
    context["holidays"] = context["page_obj"]
    return render(request, "booking/staff_holidays.html", context)


@staff_required
@require_POST
@transaction.atomic
def holiday_delete(request, holiday_id):
    item = get_object_or_404(CompanyHoliday, pk=holiday_id)
    AuditEvent.objects.create(
        actor=request.user,
        action="holiday_removed",
        target_type="holiday",
        target_id=item.pk,
        outcome="success",
    )
    item.delete()
    return redirect("staff-holidays")


@staff_required
@require_GET
def users(request):
    return _user_directory(request)


def _user_directory(request, form=None):
    qs = User.objects.order_by("email")
    search = request.GET.get("q", "").strip()[:100]
    access = request.GET.get("access", "")
    if search:
        qs = qs.filter(
            Q(email__icontains=search)
            | Q(first_name__icontains=search)
            | Q(last_name__icontains=search)
            | Q(department__icontains=search)
        )
    if access == "staff":
        qs = qs.filter(is_staff=True)
    elif access == "inactive":
        qs = qs.filter(is_active=False)
    initial = {"is_active": True}
    edit_id = request.GET.get("edit", "")[:10]
    if edit_id.isascii() and edit_id.isdigit():
        item = get_object_or_404(User, pk=int(edit_id))
        initial = {field: getattr(item, field) for field in UserAccessForm.base_fields}
    context = paginate(request, qs)
    context.update(
        {
            "users": context["page_obj"],
            "search": search,
            "access": access,
            "form": form if form is not None else UserAccessForm(initial=initial),
        }
    )
    return render(request, "booking/staff_users.html", context)


@staff_required
@require_POST
def user_save(request):
    form = UserAccessForm(request.POST)
    if not form.is_valid():
        return _user_directory(request, form)
    data = form.cleaned_data
    email = data["email"]
    with transaction.atomic():
        user, _ = User.objects.get_or_create(email__iexact=email, defaults={"email": email, "password": "!"})
        user = User.objects.select_for_update(no_key=True).get(pk=user.pk)
        active, staff = data["is_active"], data["is_staff"]
        if user.pk == request.user.pk and (not active or not staff):
            messages.error(request, "You cannot remove your own staff access")
            return redirect("staff-users")
        changed_access = (user.is_active, user.is_staff) != (active, staff)
        user.is_active, user.is_staff = active, staff
        user.first_name, user.last_name, user.department = (
            data["first_name"],
            data["last_name"],
            data["department"],
        )
        if changed_access:
            user.auth_version += 1
            EmailToken.objects.filter(
                user=user,
                purpose__in=[EmailToken.Purpose.LOGIN, EmailToken.Purpose.STAFF],
                consumed_at__isnull=True,
            ).update(consumed_at=timezone.now())
        if not staff:
            user.is_superuser = False
        user.save(
            update_fields=[
                "is_active",
                "is_staff",
                "is_superuser",
                "auth_version",
                "first_name",
                "last_name",
                "department",
            ]
        )
        AuditEvent.objects.create(
            actor=request.user,
            action="user_access_modified",
            target_type="user",
            target_id=user.pk,
            outcome="success",
            details={"active": active, "staff": staff},
        )
    messages.success(request, "User saved")
    return redirect("staff-users")


@staff_required
@require_POST
@sensitive_variables("token", "link")
def send_staff_setup(request, user_id):
    user = get_object_or_404(User, pk=user_id, is_staff=True, is_active=True)
    uid = urlsafe_base64_encode(str(user.pk).encode())
    token = default_token_generator.make_token(user)
    link = settings.PUBLIC_BASE_URL.rstrip("/") + f"/staff/set-password/{uid}/{token}/"
    try:
        deliver_mail(
            "Set your WDN staff password",
            f"Use this link to set your staff password: {link}\nIf you did not expect this, contact IT.",
            [user.email],
        )
    except Exception:
        AuditEvent.objects.create(
            actor=request.user,
            action="staff_setup_email_failed",
            target_type="user",
            target_id=user.pk,
            outcome="failed",
        )
        messages.error(request, "The setup email could not be sent")
    else:
        AuditEvent.objects.create(
            actor=request.user,
            action="staff_setup_sent",
            target_type="user",
            target_id=user.pk,
            outcome="success",
        )
        messages.success(request, "Password setup link sent")
    return redirect("staff-users")


@never_cache
@sensitive_post_parameters("password", "confirm")
@sensitive_variables("password", "token")
@require_http_methods(["GET", "POST"])
def set_staff_password(request, uid, token):
    if len(uid) > 32 or len(token) > 128:
        raise Http404
    try:
        decoded = force_str(urlsafe_base64_decode(uid))
        if not decoded.isascii() or not decoded.isdigit() or len(decoded) > 19:
            raise ValueError
        user = User.objects.get(pk=int(decoded), is_staff=True, is_active=True)
    except (ValueError, TypeError, OverflowError, UnicodeDecodeError, User.DoesNotExist):
        raise Http404 from None
    if not default_token_generator.check_token(user, token):
        raise Http404
    if request.method == "POST":
        password = request.POST.get("password", "")
        try:
            if password != request.POST.get("confirm", "") or len(password) < 12:
                raise ValidationError("Use matching passwords of at least 12 characters")
            validate_password(password, user)
        except ValidationError as exc:
            messages.error(request, " ".join(exc.messages))
        else:
            with transaction.atomic():
                user = User.objects.select_for_update().get(pk=user.pk)
                if (
                    not user.is_active
                    or not user.is_staff
                    or not default_token_generator.check_token(user, token)
                ):
                    raise Http404
                user.set_password(password)
                user.auth_version += 1
                user.save(update_fields=["password", "auth_version"])
                EmailToken.objects.filter(
                    user=user, purpose=EmailToken.Purpose.STAFF, consumed_at__isnull=True
                ).update(consumed_at=timezone.now())
                AuditEvent.objects.create(
                    actor=user,
                    action="staff_password_set",
                    target_type="user",
                    target_id=user.pk,
                    outcome="success",
                )
            messages.success(request, "Password set. Sign in with your password and email code.")
            return redirect("staff-login")
    return render(request, "booking/staff_password.html")


@staff_required
@require_GET
def audit(request):
    events = AuditEvent.objects.select_related("actor").order_by("-created_at", "-pk")
    action = request.GET.get("action", "")[:100]
    actor = request.GET.get("actor", "").strip()[:100]
    if action:
        events = events.filter(action=action)
    if actor:
        events = events.filter(actor__email__icontains=actor)
    context = paginate(request, events, 50)
    context.update(
        {
            "events": context["page_obj"],
            "action": action,
            "actor_search": actor,
            "actions": AuditEvent.objects.order_by("action").values_list("action", flat=True).distinct(),
        }
    )
    return render(request, "booking/staff_audit.html", context)


def _range(request):
    today = timezone.localdate()
    try:
        start = date.fromisoformat(request.GET.get("from", ""))
    except ValueError:
        start = today.replace(day=1)
    try:
        end = date.fromisoformat(request.GET.get("to", ""))
    except ValueError:
        end = today
    if end < start or (end - start).days > 366 or end == date.max:
        raise ValueError("Choose a valid range of up to one year")
    return start, end


def _report_data(start, end):
    return report_data(start, end)


@staff_required
@require_GET
def reports(request):
    try:
        start, end = _range(request)
    except ValueError as exc:
        messages.error(request, str(exc))
        start, end = timezone.localdate(), timezone.localdate()
    records, room_rows, departments, organizers = _report_data(start, end)
    today = timezone.localdate()
    quarter_start = today.replace(month=((today.month - 1) // 3) * 3 + 1, day=1)
    presets = [
        ("Today", today, today),
        ("This week", today - timedelta(days=today.weekday()), today),
        ("This month", today.replace(day=1), today),
        ("This quarter", quarter_start, today),
        ("This year", today.replace(month=1, day=1), today),
    ]
    return render(
        request,
        "booking/staff_reports.html",
        {
            "start": start,
            "end": end,
            "records": records,
            "room_rows": room_rows,
            "departments": departments,
            "organizers": organizers,
            "presets": presets,
            "total": len(records),
        },
    )


@staff_required
@require_GET
def report_excel(request):
    try:
        start, end = _range(request)
    except ValueError:
        raise Http404 from None
    records, room_rows, departments, organizers = _report_data(start, end)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bookings"
    sheet.append(["ID", "Date", "Start", "End", "Room", "Organizer", "Department", "Status", "Title"])
    for item in records:
        local_start, local_end = timezone.localtime(item.starts_at), timezone.localtime(item.ends_at)
        sheet.append(
            [
                item.pk,
                local_start.date().isoformat(),
                local_start.strftime("%H:%M"),
                local_end.strftime("%H:%M"),
                item.room.name,
                item.organizer.email,
                item.department,
                item.status,
                item.title,
            ]
        )
    room_sheet = workbook.create_sheet("Room usage")
    room_sheet.append(["Room", "Bookings", "Completed", "Cancelled", "No-shows", "Utilization %"])
    for item in room_rows:
        room_sheet.append(
            [
                item["room"].name,
                item["bookings"],
                item["completed"],
                item["cancelled"],
                item["no_shows"],
                item["utilization"],
            ]
        )
    dept_sheet = workbook.create_sheet("Departments")
    dept_sheet.append(["Department", "Bookings"])
    for department, count in departments:
        dept_sheet.append([department, count])
    organizer_sheet = workbook.create_sheet("Organizers")
    organizer_sheet.append(["Organizer", "Bookings"])
    for email, count in organizers:
        organizer_sheet.append([email, count])
    for tab in workbook:
        # User-controlled spreadsheet values must be literal strings, never formulas.
        for row in tab:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="104068")
        for column in tab.columns:
            letter = column[0].column_letter
            tab.column_dimensions[letter].width = min(
                50, max(13, max(len(str(cell.value or "")) for cell in column) + 2)
            )
    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="meeting-room-report-{start}-{end}.xlsx"'
    return response
