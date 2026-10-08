import io
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, time, timedelta
from threading import Barrier, Event
from unittest.mock import patch
from zoneinfo import ZoneInfo

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection, connections
from django.test import Client, SimpleTestCase, TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from booking.models import AuditEvent, BookingPolicy, CompanyHoliday, Notification, Reservation, Room, User
from booking.services.auth_security import staff_session_verified
from booking.services.bookings import BookingError, create_booking
from booking.services.holiday_import import (
    MAX_EXPANDED_BYTES,
    MAX_IMPORT_ROWS,
    MAX_UPLOAD_BYTES,
    HolidayImportChanged,
    HolidayImportError,
    apply_holiday_import,
    build_holiday_preview,
    parse_holiday_workbook,
)

LOCAL_TZ = ZoneInfo("Asia/Kathmandu")
NOW = datetime(2026, 10, 8, 8, tzinfo=LOCAL_TZ)
DAY = date(2026, 10, 9)
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_upload(
    rows=(), *, headers=("Date", "Occasion"), sheet="Holidays", epoch=None, filename="holidays.xlsx"
):
    workbook = Workbook()
    workbook.active.title = sheet
    if epoch is not None:
        workbook.epoch = epoch
    workbook.active.append(list(headers))
    for row in rows:
        workbook.active.append(list(row))
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile(filename, output.getvalue(), content_type=XLSX_TYPE)


class HolidayWorkbookParserTests(SimpleTestCase):
    def parse(self, rows=(), **options):
        return parse_holiday_workbook(workbook_upload(rows, **options))

    def test_excel_dates_and_strict_iso_text_are_normalized_with_source_row_numbers(self):
        result = self.parse(
            [
                (date(2026, 10, 9), "  Company holiday  "),
                (None, None),
                (datetime(2026, 10, 12), "Festival"),
                ("2026-10-13", "  Second festival day  "),
            ]
        )
        self.assertEqual(result["errors"], [])
        self.assertEqual(
            result["rows"],
            [
                {"row_number": 2, "date": "2026-10-09", "name": "Company holiday"},
                {"row_number": 4, "date": "2026-10-12", "name": "Festival"},
                {"row_number": 5, "date": "2026-10-13", "name": "Second festival day"},
            ],
        )

    def test_headers_ignore_case_and_surrounding_whitespace(self):
        result = self.parse([("2026-10-09", "Holiday")], headers=(" date ", " OCCASION "))
        self.assertEqual(result["errors"], [])
        self.assertEqual(len(result["rows"]), 1)

    def test_date_conversion_respects_workbooks_using_the_1904_epoch(self):
        result = self.parse([(DAY, "Holiday")], epoch=CALENDAR_MAC_1904)
        self.assertEqual(result["errors"], [])
        self.assertEqual(result["rows"][0]["date"], DAY.isoformat())

    def test_holiday_names_allow_160_characters_after_trimming(self):
        result = self.parse([(DAY, "  " + "x" * 160 + "  ")])
        self.assertEqual(result["errors"], [])
        self.assertEqual(result["rows"][0]["name"], "x" * 160)
        result = self.parse([(DAY, "x" * 161)])
        self.assertTrue(result["errors"])
        self.assertEqual(result["errors"][0]["row_number"], 2)

    def test_missing_sheet_bad_headers_and_extra_columns_are_rejected(self):
        for options in (
            {"sheet": "Sheet1"},
            {"headers": ("Date", "Name")},
            {"headers": ("Occasion", "Date")},
            {"headers": ("Date", "Occasion", "Notes")},
        ):
            with self.subTest(options=options):
                self.assertTrue(self.parse([(DAY, "Holiday")], **options)["errors"])
        self.assertTrue(self.parse([(DAY, "Holiday", "Unexpected cell")])["errors"])

    def test_empty_workbook_and_blank_rows_do_not_create_phantom_holidays(self):
        for rows in ([], [(None, None)], [("   ", "  "), (None, None)]):
            with self.subTest(rows=rows):
                result = self.parse(rows)
                self.assertTrue(result["errors"])
                self.assertEqual(result["rows"], [])

    def test_blank_name_missing_date_and_non_text_names_report_row_errors(self):
        for row in ((DAY, None), (DAY, "  "), (None, "Holiday"), (DAY, 123), (DAY, True)):
            with self.subTest(row=row):
                result = self.parse([row])
                self.assertTrue(result["errors"])
                self.assertEqual(result["errors"][0]["row_number"], 2)

    def test_invalid_ambiguous_numeric_and_time_component_dates_are_rejected(self):
        for value in (
            "09/10/2026",
            "2026-10-9",
            "2026-02-30",
            "2083-13-01",
            "2026-10-09T00:00:00",
            46304,
            True,
            datetime(2026, 10, 9, 9, 30),
        ):
            with self.subTest(value=value):
                result = self.parse([(value, "Holiday")])
                self.assertTrue(result["errors"])
                self.assertEqual(result["errors"][0]["row_number"], 2)

    def test_formulas_in_dates_names_and_headers_are_rejected_instead_of_evaluated(self):
        for row in (("=DATE(2026,10,9)", "Holiday"), (DAY, '=HYPERLINK("https://example.com","Holiday")')):
            with self.subTest(row=row):
                self.assertTrue(self.parse([row])["errors"])
        self.assertTrue(self.parse([(DAY, "Holiday")], headers=("=1+1", "Occasion"))["errors"])

    def test_duplicate_dates_in_one_workbook_are_reported_on_the_duplicate_row(self):
        result = self.parse([(DAY, "First holiday"), (DAY.isoformat(), "Another name")])
        self.assertTrue(result["errors"])
        self.assertTrue(any(error["row_number"] == 3 for error in result["errors"]))
        self.assertTrue(any("duplicate" in error["message"].lower() for error in result["errors"]))

    def test_non_xlsx_extension_and_corrupt_archive_are_rejected(self):
        for filename in ("holidays.xls", "holidays.csv", "holidays.xlsm"):
            with self.subTest(filename=filename):
                self.assertTrue(self.parse([(DAY, "Holiday")], filename=filename)["errors"])
        for content in (b"", b"not an Excel workbook", b"PK\x03\x04broken"):
            result = parse_holiday_workbook(
                SimpleUploadedFile("holidays.xlsx", content, content_type=XLSX_TYPE)
            )
            self.assertTrue(result["errors"])
            self.assertEqual(result["rows"], [])

    def test_oversized_upload_is_rejected_before_openpyxl_reads_it(self):
        upload = SimpleUploadedFile("holidays.xlsx", b"x" * (MAX_UPLOAD_BYTES + 1), content_type=XLSX_TYPE)
        with patch("booking.services.holiday_import.load_workbook") as load:
            result = parse_holiday_workbook(upload)
        self.assertTrue(result["errors"])
        load.assert_not_called()

    def test_expanded_zip_size_is_bounded_before_workbook_parsing(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/worksheets/sheet1.xml", b"x" * (MAX_EXPANDED_BYTES + 1))
        upload = SimpleUploadedFile("holidays.xlsx", output.getvalue(), content_type=XLSX_TYPE)
        self.assertLess(upload.size, MAX_UPLOAD_BYTES)
        with patch("booking.services.holiday_import.load_workbook") as load:
            result = parse_holiday_workbook(upload)
        self.assertTrue(result["errors"])
        load.assert_not_called()

    def test_row_limit_rejects_entire_upload(self):
        first = date(2026, 1, 1)
        result = self.parse(
            [(first + timedelta(days=index), "Holiday") for index in range(MAX_IMPORT_ROWS + 1)]
        )
        self.assertTrue(result["errors"])

    def test_macros_and_xml_entity_declarations_are_rejected_before_parsing(self):
        for filename, contents in (
            ("xl/vbaProject.bin", b"macro content"),
            ("xl/workbook.xml", b'<!DOCTYPE x [<!ENTITY data SYSTEM "file:///etc/passwd">]><x>&data;</x>'),
        ):
            with self.subTest(filename=filename):
                output = io.BytesIO()
                with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                    archive.writestr(filename, contents)
                upload = SimpleUploadedFile("holidays.xlsx", output.getvalue(), content_type=XLSX_TYPE)
                with patch("booking.services.holiday_import.load_workbook") as load:
                    result = parse_holiday_workbook(upload)
                self.assertTrue(result["errors"])
                load.assert_not_called()


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class HolidayImportTests(TestCase):
    def setUp(self):
        self.clock = patch("django.utils.timezone.now", return_value=NOW)
        self.clock.start()
        self.addCleanup(self.clock.stop)
        BookingPolicy.objects.get_or_create(pk=1)
        self.staff = User.objects.create_user("holiday.desk@wdn.com.np", is_staff=True)
        self.other_staff = User.objects.create_user("holiday.admin@transgate.com.np", is_staff=True)
        self.employee = User.objects.create_user("holiday.employee@wdn.com.np", department="Accounts")
        self.room = Room.objects.create(name="Holiday room", location="WDN", floor="2", capacity=8)
        self.existing = CompanyHoliday.objects.create(date=date(2026, 10, 12), name="Old festival name")
        self.untouched = CompanyHoliday.objects.create(
            date=date(2026, 12, 25), name="Existing company holiday"
        )
        self.import_url = reverse("staff-holiday-import")
        self.template_url = reverse("staff-holiday-template")
        self.list_url = reverse("staff-holidays")
        self.login(self.client, self.staff)

    def login(self, client, user, *, verified=True, version=None):
        client.force_login(user)
        session = client.session
        session["staff_verified"] = verified
        session["staff_auth_version"] = user.auth_version if version is None else version
        session.save()

    def preview(self, rows=None, *, client=None, **options):
        rows = [(DAY, "Company celebration")] if rows is None else rows
        return (client or self.client).post(
            self.import_url, {"action": "preview", "file": workbook_upload(rows, **options)}
        )

    def confirm(self, preview, *, client=None, acknowledge=True, **extra):
        values = {"action": "confirm", "preview_id": preview.context["confirm_form"]["preview_id"].value()}
        if acknowledge:
            values["acknowledge_bookings"] = "on"
        values.update(extra)
        return (client or self.client).post(self.import_url, values)

    def meeting(self, *, day=DAY, status="approved", starts_hour=11, duration=60, room=None):
        start = datetime.combine(day, time(starts_hour), LOCAL_TZ)
        return Reservation.objects.create(
            room=room or self.room,
            organizer=self.employee,
            created_by=self.employee,
            kind="booking",
            status=status,
            title="Existing meeting on imported holiday",
            department="Accounts",
            starts_at=start,
            ends_at=start + timedelta(minutes=duration),
            occupied_from=start,
            occupied_until=start + timedelta(minutes=duration + 15),
        )

    def test_holiday_list_links_upload_and_blank_template_and_get_does_not_import(self):
        response = self.client.get(self.list_url)
        self.assertContains(response, self.import_url)
        self.assertContains(response, self.template_url)
        response = self.client.get(self.import_url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="csrfmiddlewaretoken"')
        self.assertContains(response, 'type="file"')
        self.assertEqual(CompanyHoliday.objects.count(), 2)

    def test_template_download_is_header_only_xlsx_with_holidays_sheet(self):
        response = self.client.get(self.template_url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX_TYPE)
        self.assertIn("MBS_Holiday_Template.xlsx", response["Content-Disposition"])
        workbook = load_workbook(io.BytesIO(response.content), data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Holidays"])
        rows = list(workbook["Holidays"].values)
        self.assertEqual(rows[0][:2], ("Date", "Occasion"))
        self.assertFalse(any(any(value is not None for value in row) for row in rows[1:]))
        self.assertEqual(CompanyHoliday.objects.count(), 2)
        self.assertFalse(AuditEvent.objects.filter(action__startswith="holiday_").exists())

    def test_preview_shows_new_update_unchanged_actions_without_mutating_database(self):
        response = self.preview(
            [
                (DAY, "Company celebration"),
                (self.existing.date, "Updated festival"),
                (self.untouched.date, self.untouched.name),
            ]
        )
        self.assertEqual(response.status_code, 200)
        preview = response.context["preview"]
        self.assertEqual({row["action"] for row in preview["rows"]}, {"new", "update", "unchanged"})
        self.assertEqual(preview["summary"]["new"], 1)
        self.assertEqual(preview["summary"]["update"], 1)
        self.assertEqual(preview["summary"]["unchanged"], 1)
        updated = next(row for row in preview["rows"] if row["action"] == "update")
        self.assertEqual(updated["previous_name"], self.existing.name)
        self.assertTrue(response.context["confirm_form"]["preview_id"].value())
        self.assertEqual(CompanyHoliday.objects.count(), 2)
        self.existing.refresh_from_db()
        self.assertEqual(self.existing.name, "Old festival name")
        self.assertFalse(AuditEvent.objects.filter(action__startswith="holiday_").exists())

    def test_confirm_upserts_preview_rows_and_keeps_holidays_not_in_workbook(self):
        preview = self.preview([(DAY, "Company celebration"), (self.existing.date, "Updated festival")])
        response = self.confirm(preview)
        self.assertRedirects(response, self.list_url)
        self.assertEqual(CompanyHoliday.objects.get(date=DAY).name, "Company celebration")
        self.existing.refresh_from_db()
        self.assertEqual(self.existing.name, "Updated festival")
        self.untouched.refresh_from_db()
        self.assertEqual(self.untouched.name, "Existing company holiday")
        self.assertEqual(CompanyHoliday.objects.count(), 3)
        events = AuditEvent.objects.filter(action__startswith="holiday_")
        self.assertTrue(events.exists())
        self.assertTrue(all(event.actor == self.staff and event.outcome == "success" for event in events))

    def test_confirm_uses_server_preview_and_ignores_forged_posted_rows(self):
        preview = self.preview()
        response = self.confirm(
            preview, rows='[{"date":"2026-10-12","name":"Forged name"}]', name="Forged name"
        )
        self.assertRedirects(response, self.list_url)
        self.assertEqual(CompanyHoliday.objects.get(date=DAY).name, "Company celebration")
        self.existing.refresh_from_db()
        self.assertEqual(self.existing.name, "Old festival name")

    def test_any_invalid_row_prevents_partial_import_and_reports_its_row_number(self):
        response = self.preview([(DAY, "Valid company holiday"), ("not a date", "Invalid row")])
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["errors"])
        self.assertTrue(any(error["row_number"] == 3 for error in response.context["errors"]))
        self.assertFalse(response.context.get("confirm_form"))
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())
        self.assertFalse(AuditEvent.objects.filter(action__startswith="holiday_").exists())

    def test_wrong_extension_and_corrupt_upload_return_friendly_errors_without_writes(self):
        for upload in (
            workbook_upload([(DAY, "Holiday")], filename="holidays.xls"),
            SimpleUploadedFile("holidays.xlsx", b"broken workbook", content_type=XLSX_TYPE),
        ):
            with self.subTest(filename=upload.name):
                response = self.client.post(self.import_url, {"action": "preview", "file": upload})
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context["errors"])
        self.assertEqual(CompanyHoliday.objects.count(), 2)

    def test_preview_counts_active_meetings_and_requires_acknowledgment_before_import(self):
        approved = self.meeting()
        pending = self.meeting(starts_hour=14, status="pending")
        cancelled = self.meeting(starts_hour=16, status="cancelled", duration=30)
        before = list(Reservation.objects.order_by("pk").values())
        preview = self.preview()
        self.assertEqual(preview.context["preview"]["rows"][0]["booking_count"], 2)
        self.assertEqual(preview.context["preview"]["summary"]["affected_bookings"], 2)
        response = self.confirm(preview, acknowledge=False)
        self.assertEqual(response.status_code, 200)
        self.assertIn("acknowledge_bookings", response.context["confirm_form"].errors)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())
        self.assertRedirects(self.confirm(response), self.list_url)
        self.assertEqual(list(Reservation.objects.order_by("pk").values()), before)
        self.assertFalse(Notification.objects.exists())
        self.assertEqual(Reservation.objects.get(pk=approved.pk).status, "approved")
        self.assertEqual(Reservation.objects.get(pk=pending.pk).status, "pending")
        self.assertEqual(Reservation.objects.get(pk=cancelled.pk).status, "cancelled")

    def test_no_existing_meetings_import_does_not_require_the_acknowledgment_checkbox(self):
        response = self.confirm(self.preview(), acknowledge=False)
        self.assertRedirects(response, self.list_url)
        self.assertTrue(CompanyHoliday.objects.filter(date=DAY).exists())

    def test_imported_holidays_close_calendar_and_employee_booking_availability(self):
        self.assertRedirects(self.confirm(self.preview()), self.list_url)
        self.client.force_login(self.employee)
        month = self.client.get(reverse("calendar"), {"date": DAY.isoformat(), "view": "month"})
        days = {day["date"]: day for week in month.context["calendar"]["month_weeks"] for day in week["days"]}
        self.assertIn("Company celebration", days[DAY]["closed_reason"])
        self.assertFalse(days[DAY]["new_booking_url"])
        day = self.client.get(reverse("calendar"), {"date": DAY.isoformat(), "view": "day"})
        self.assertTrue(day.context["slots"])
        self.assertTrue(all(not cell["available"] for slot in day.context["slots"] for cell in slot["cells"]))
        response = self.client.get(
            reverse("booking-availability"),
            {"date": DAY.isoformat(), "start_time": "11:00", "end_time": "12:00", "recurrence": "none"},
        )
        self.assertFalse(response.json()["ok"])
        self.assertIn("working day", response.json()["error"])

    def test_preview_and_template_require_active_staff_with_current_mfa_verification(self):
        clients = [Client()]
        for user, verified, version in (
            (self.employee, True, None),
            (self.staff, False, None),
            (self.staff, True, self.staff.auth_version + 1),
        ):
            client = Client()
            self.login(client, user, verified=verified, version=version)
            clients.append(client)
        for client in clients:
            for url in (self.import_url, self.template_url):
                with self.subTest(url=url, client=client):
                    self.assertRedirects(client.get(url), reverse("staff-login"))
            self.assertRedirects(self.preview(client=client), reverse("staff-login"))
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())

    def test_mutations_require_csrf_for_preview_and_confirmation(self):
        client = Client(enforce_csrf_checks=True)
        self.login(client, self.staff)
        self.assertEqual(self.preview(client=client).status_code, 403)
        page = client.get(self.import_url)
        token = str(page.context["csrf_token"])
        preview = client.post(
            self.import_url,
            {"action": "preview", "file": workbook_upload([(DAY, "Holiday")]), "csrfmiddlewaretoken": token},
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(self.confirm(preview, client=client).status_code, 403)
        self.assertRedirects(self.confirm(preview, client=client, csrfmiddlewaretoken=token), self.list_url)

    def test_template_only_accepts_get_and_import_rejects_other_methods(self):
        self.assertEqual(self.client.post(self.template_url).status_code, 405)
        self.assertEqual(self.client.put(self.import_url).status_code, 405)

    def test_missing_preview_and_replayed_confirmation_cannot_write(self):
        response = self.client.post(self.import_url, {"action": "confirm", "preview_id": "not-real"})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())
        preview = self.preview()
        self.assertRedirects(self.confirm(preview), self.list_url)
        audit_count = AuditEvent.objects.count()
        self.assertEqual(self.confirm(preview).status_code, 400)
        self.assertEqual(AuditEvent.objects.count(), audit_count)

    def test_malformed_unicode_confirmation_id_returns_validation_error(self):
        self.preview()
        response = self.client.post(self.import_url, {"action": "confirm", "preview_id": "नक्कली-preview"})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())

    def test_preview_expires_after_fifteen_minutes(self):
        preview = self.preview()
        with patch("django.utils.timezone.now", return_value=NOW + timedelta(minutes=16)):
            response = self.confirm(preview)
        self.assertEqual(response.status_code, 400)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())

    def test_another_staff_session_cannot_confirm_someone_elses_preview(self):
        preview = self.preview()
        client = Client()
        self.login(client, self.other_staff)
        self.assertEqual(self.confirm(preview, client=client).status_code, 400)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())

    def test_logout_discards_preview_even_after_same_staff_signs_in_again(self):
        preview = self.preview()
        self.client.logout()
        self.login(self.client, self.staff)
        self.assertEqual(self.confirm(preview).status_code, 400)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())

    def test_revoked_or_version_changed_staff_cannot_confirm_existing_preview(self):
        for change in ({"is_staff": False}, {"is_active": False}, {"auth_version": 2}):
            with self.subTest(change=change):
                User.objects.filter(pk=self.staff.pk).update(is_staff=True, is_active=True, auth_version=1)
                self.staff.refresh_from_db()
                self.login(self.client, self.staff)
                preview = self.preview()
                User.objects.filter(pk=self.staff.pk).update(**change)
                self.assertRedirects(self.confirm(preview), reverse("staff-login"))
                self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())

    def test_revocation_between_initial_check_and_import_is_rechecked_before_mutation(self):
        preview = self.preview()

        def revoke_after_check(request):
            allowed = staff_session_verified(request)
            User.objects.filter(pk=self.staff.pk).update(is_staff=False)
            return allowed

        with patch("booking.staff_views.staff_session_verified", side_effect=revoke_after_check):
            response = self.confirm(preview)
        self.assertRedirects(response, reverse("staff-login"))
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())
        self.assertFalse(AuditEvent.objects.filter(action__startswith="holiday_").exists())

    def test_changed_existing_holiday_refreshes_preview_instead_of_overwriting_concurrent_edit(self):
        preview = self.preview([(DAY, "New holiday"), (self.existing.date, "Imported name")])
        old_id = preview.context["confirm_form"]["preview_id"].value()
        CompanyHoliday.objects.filter(pk=self.existing.pk).update(name="Other staff edit")
        response = self.confirm(preview)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())
        self.existing.refresh_from_db()
        self.assertEqual(self.existing.name, "Other staff edit")
        self.assertNotEqual(response.context["confirm_form"]["preview_id"].value(), old_id)
        row = next(
            row
            for row in response.context["preview"]["rows"]
            if row["date"] == self.existing.date.isoformat()
        )
        self.assertEqual(row["previous_name"], "Other staff edit")
        self.assertRedirects(self.confirm(response), self.list_url)
        self.existing.refresh_from_db()
        self.assertEqual(self.existing.name, "Imported name")

    def test_new_meeting_after_preview_requires_a_fresh_review_and_acknowledgment(self):
        preview = self.preview()
        booking = self.meeting()
        response = self.confirm(preview)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())
        self.assertEqual(response.context["preview"]["summary"]["affected_bookings"], 1)
        missing_ack = self.confirm(response, acknowledge=False)
        self.assertIn("acknowledge_bookings", missing_ack.context["confirm_form"].errors)
        self.assertRedirects(self.confirm(missing_ack), self.list_url)
        booking.refresh_from_db()
        self.assertEqual(booking.status, "approved")

    def test_replacing_preview_invalidates_older_confirmation_id(self):
        old = self.preview([(DAY, "Old uploaded holiday")])
        current = self.preview([(self.existing.date, "Current uploaded holiday")])
        self.assertEqual(self.confirm(old).status_code, 400)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())
        self.assertRedirects(self.confirm(current), self.list_url)
        self.existing.refresh_from_db()
        self.assertEqual(self.existing.name, "Current uploaded holiday")

    def test_discard_preview_returns_to_upload_form_without_importing(self):
        preview = self.preview()
        response = self.client.post(self.import_url, {"action": "discard"})
        self.assertRedirects(response, self.import_url)
        page = self.client.get(self.import_url)
        self.assertIsNone(page.context["preview"])
        self.assertTrue(page.context["upload_form"])
        self.assertContains(page, 'type="file"')
        self.assertEqual(self.confirm(preview).status_code, 400)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())

    def test_cancel_preview_with_existing_meetings_does_not_require_acknowledgment(self):
        booking = self.meeting()
        preview = self.preview()
        self.assertEqual(preview.context["preview"]["summary"]["affected_bookings"], 1)
        response = self.client.post(self.import_url, {"action": "cancel"})
        self.assertRedirects(response, self.list_url)
        page = self.client.get(self.import_url)
        self.assertIsNone(page.context["preview"])
        self.assertEqual(self.confirm(preview).status_code, 400)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())
        booking.refresh_from_db()
        self.assertEqual(booking.status, "approved")
        self.assertFalse(Notification.objects.exists())
        self.assertFalse(AuditEvent.objects.filter(action__startswith="holiday_").exists())

    def test_holiday_names_are_escaped_in_preview_and_calendar(self):
        name = '<script>alert("holiday")</script>'
        preview = self.preview([(DAY, name)])
        self.assertContains(preview, "&lt;script&gt;")
        self.assertNotContains(preview, name)
        self.assertRedirects(self.confirm(preview), self.list_url)
        self.client.force_login(self.employee)
        calendar = self.client.get(reverse("calendar"), {"date": DAY.isoformat()})
        self.assertContains(calendar, "&lt;script&gt;")
        self.assertNotContains(calendar, name)

    def test_preview_counts_cross_midnight_meeting_once_in_summary_and_on_each_affected_day(self):
        self.meeting(starts_hour=23, duration=120)
        preview = self.preview([(DAY, "First day"), (DAY + timedelta(days=1), "Second day")])
        self.assertEqual([row["booking_count"] for row in preview.context["preview"]["rows"]], [1, 1])
        self.assertEqual(preview.context["preview"]["summary"]["affected_bookings"], 1)

    def test_meeting_ending_at_midnight_does_not_count_its_buffer_as_a_holiday_meeting(self):
        self.meeting(day=DAY - timedelta(days=1), starts_hour=23, duration=60)
        preview = self.preview()
        self.assertEqual(preview.context["preview"]["rows"][0]["booking_count"], 0)
        self.assertEqual(preview.context["preview"]["summary"]["affected_bookings"], 0)

    def test_audit_failure_rolls_back_all_holiday_changes_and_earlier_audit_rows(self):
        rows = parse_holiday_workbook(
            workbook_upload([(DAY, "New holiday"), (self.existing.date, "Changed festival")])
        )["rows"]
        preview = build_holiday_preview(rows)
        create_audit = AuditEvent.objects.create
        audit_calls = 0

        def fail_second_audit(**values):
            nonlocal audit_calls
            audit_calls += 1
            if audit_calls == 2:
                raise RuntimeError("Audit storage unavailable")
            return create_audit(**values)

        with patch(
            "booking.services.holiday_import.AuditEvent.objects.create", side_effect=fail_second_audit
        ):
            with self.assertRaises(RuntimeError):
                apply_holiday_import(rows, self.staff, preview["snapshot"])
        self.assertEqual(audit_calls, 2)
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())
        self.existing.refresh_from_db()
        self.assertEqual(self.existing.name, "Old festival name")
        self.assertFalse(AuditEvent.objects.filter(action__startswith="holiday_").exists())

    def test_service_requires_staff_and_matching_preview_snapshot_before_writes(self):
        rows = parse_holiday_workbook(workbook_upload([(DAY, "New holiday")]))["rows"]
        preview = build_holiday_preview(rows)
        with self.assertRaises(HolidayImportError):
            apply_holiday_import(rows, self.employee, preview["snapshot"])
        with self.assertRaises(HolidayImportChanged):
            apply_holiday_import(rows, self.staff, "outdated-snapshot")
        self.assertFalse(CompanyHoliday.objects.filter(date=DAY).exists())
        self.assertFalse(AuditEvent.objects.filter(action__startswith="holiday_").exists())


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class HolidayImportConcurrencyTests(TransactionTestCase):
    def test_booking_waits_for_import_commit_then_rejects_new_holiday(self):
        clock = patch("django.utils.timezone.now", return_value=NOW)
        clock.start()
        self.addCleanup(clock.stop)
        BookingPolicy.objects.get_or_create(pk=1)
        staff = User.objects.create_user("import.race.desk@wdn.com.np", is_staff=True)
        employee = User.objects.create_user("import.race.employee@wdn.com.np")
        room = Room.objects.create(name="Import race room", location="WDN", floor="2", capacity=8)
        rows = parse_holiday_workbook(workbook_upload([(DAY, "Newly closed date")]))["rows"]
        snapshot = build_holiday_preview(rows)["snapshot"]
        import_has_rooms, allow_import_commit = Event(), Event()
        booking_attempted_room, booking_has_room = Event(), Event()
        data = {
            "room": room,
            "date": DAY,
            "start_time": time(11),
            "end_time": time(12),
            "title": "Meeting racing with holiday import",
            "description": "",
            "meeting_type": "internal",
            "guest_company_name": "",
            "external_attendee_count": 0,
            "department": "Accounts",
            "attendees": [],
            "refreshments_requested": False,
            "front_desk_notes": "",
            "recurrence": "none",
            "until_date": None,
        }

        def paused_final_preview(import_rows):
            import_has_rooms.set()
            if not allow_import_commit.wait(timeout=10):
                raise TimeoutError("Import was not released")
            return build_holiday_preview(import_rows)

        def import_holiday():
            connections.close_all()
            try:
                return apply_holiday_import(rows, staff, snapshot)
            finally:
                connections.close_all()

        def observe_room_lock(execute, sql, params, many, context):
            room_lock = 'FROM "rooms"' in sql and "FOR UPDATE" in sql.upper()
            if room_lock:
                booking_attempted_room.set()
            result = execute(sql, params, many, context)
            if room_lock:
                booking_has_room.set()
            return result

        def book_meeting():
            connections.close_all()
            try:
                with connection.execute_wrapper(observe_room_lock):
                    try:
                        create_booking(data, actor=employee)
                    except BookingError as error:
                        return str(error)
                return "unexpectedly booked"
            finally:
                connections.close_all()

        with (
            patch("booking.services.holiday_import.build_holiday_preview", side_effect=paused_final_preview),
            ThreadPoolExecutor(max_workers=2) as executor,
        ):
            import_result = executor.submit(import_holiday)
            try:
                self.assertTrue(import_has_rooms.wait(timeout=10))
                booking_result = executor.submit(book_meeting)
                self.assertTrue(booking_attempted_room.wait(timeout=10))
                self.assertFalse(booking_has_room.wait(timeout=0.2))
            finally:
                allow_import_commit.set()
            self.assertEqual(import_result.result(timeout=10)["new"], 1)
            self.assertIn("working day", booking_result.result(timeout=10))
        self.assertTrue(booking_has_room.is_set())
        self.assertTrue(CompanyHoliday.objects.filter(date=DAY, name="Newly closed date").exists())
        self.assertFalse(Reservation.objects.exists())
        self.assertFalse(Notification.objects.exists())

    def test_concurrent_confirmation_consumes_one_preview_and_imports_once(self):
        clock = patch("django.utils.timezone.now", return_value=NOW)
        clock.start()
        self.addCleanup(clock.stop)
        staff = User.objects.create_user("concurrent.holiday.desk@wdn.com.np", is_staff=True)
        client = Client()
        client.force_login(staff)
        session = client.session
        session["staff_verified"] = True
        session["staff_auth_version"] = staff.auth_version
        session.save()
        url = reverse("staff-holiday-import")
        preview = client.post(url, {"action": "preview", "file": workbook_upload([(DAY, "Company holiday")])})
        preview_id = preview.context["confirm_form"]["preview_id"].value()
        barrier = Barrier(2)

        def confirm_once(_):
            connections.close_all()
            try:
                request_client = Client()
                request_client.cookies = client.cookies.copy()
                barrier.wait(timeout=10)
                return request_client.post(url, {"action": "confirm", "preview_id": preview_id}).status_code
            finally:
                connections.close_all()

        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(confirm_once, range(2)))
        self.assertEqual(sorted(results), [302, 400])
        self.assertEqual(CompanyHoliday.objects.filter(date=DAY).count(), 1)
        self.assertEqual(AuditEvent.objects.filter(action="holiday_imported").count(), 1)
        self.assertEqual(AuditEvent.objects.filter(action="holiday_saved").count(), 1)
