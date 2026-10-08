"""Bounded Excel parsing and atomic company holiday imports."""

import hashlib
import io
import json
import re
from bisect import bisect_left, bisect_right
from datetime import date, datetime, time, timedelta
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo

from django.db import IntegrityError, transaction
from openpyxl import load_workbook

from booking.models import AuditEvent, CompanyHoliday, Reservation, Room

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_IMPORT_ROWS = 1000
LOCAL_TZ = ZoneInfo("Asia/Kathmandu")
REVIEW_STATUSES = (
    Reservation.Status.PENDING,
    Reservation.Status.APPROVED,
    Reservation.Status.CHECKED_IN,
)
ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}\Z")
FORMULA_ELEMENT = re.compile(rb"<(?:[\w.-]+:)?f(?:\s|/?>)", re.IGNORECASE)


class HolidayImportError(Exception):
    """The import cannot be safely applied."""


class HolidayImportChanged(HolidayImportError):
    """Holiday or booking data changed after the preview."""


def _error(message, row_number=None):
    return {"row_number": row_number, "message": message}


def _failure(message):
    return {"rows": [], "errors": [_error(message)]}


def _safe_archive(contents):
    """Preflight ZIP expansion and XML before handing it to openpyxl."""
    with ZipFile(io.BytesIO(contents)) as archive:
        entries = archive.infolist()
        if len(entries) > 500 or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
            raise HolidayImportError("The workbook is too large when opened. Use the holiday template.")
        expanded = 0
        for item in entries:
            name = item.filename.lower()
            if item.flag_bits & 1:
                raise HolidayImportError("Password-protected workbooks are not supported.")
            if "vbaproject" in name or name.endswith(".bin"):
                raise HolidayImportError("Macros are not supported. Upload a plain .xlsx workbook.")
            with archive.open(item) as member:
                content = member.read(MAX_EXPANDED_BYTES - expanded + 1)
            expanded += len(content)
            if expanded > MAX_EXPANDED_BYTES:
                raise HolidayImportError("The workbook is too large when opened. Use the holiday template.")
            if name.endswith(".xml"):
                # This also catches UTF-16 declarations; Excel XML must not expand entities.
                declarations = content.replace(b"\x00", b"").upper()
                if b"<!DOCTYPE" in declarations or b"<!ENTITY" in declarations:
                    raise HolidayImportError("This workbook contains unsupported XML declarations.")
                if FORMULA_ELEMENT.search(content):
                    raise HolidayImportError(
                        "Formulas are not supported. Replace formulas with their values."
                    )
                if name == "[content_types].xml" and b"MACROENABLED" in declarations:
                    raise HolidayImportError("Macros are not supported. Upload a plain .xlsx workbook.")


def _holiday_date(value):
    if isinstance(value, datetime):
        if value.time() != time.min or value.tzinfo is not None:
            raise ValueError
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        if ISO_DATE.fullmatch(value):
            return date.fromisoformat(value)
    raise ValueError


def parse_holiday_workbook(upload):
    """Return sanitized rows and row-numbered errors, without database writes."""
    if PurePath(str(getattr(upload, "name", ""))).suffix.lower() != ".xlsx":
        return _failure("Upload an Excel .xlsx file.")
    if getattr(upload, "size", 0) > MAX_UPLOAD_BYTES:
        return _failure("The workbook must be no larger than 2 MiB.")
    try:
        upload.seek(0)
        contents = upload.read(MAX_UPLOAD_BYTES + 1)
    except (AttributeError, OSError, ValueError):
        return _failure("The workbook could not be read. Select the file again.")
    if len(contents) > MAX_UPLOAD_BYTES:
        return _failure("The workbook must be no larger than 2 MiB.")
    if not contents:
        return _failure("The workbook is empty.")
    workbook = None
    try:
        _safe_archive(contents)
        workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=False, keep_links=False)
        if "Holidays" not in workbook.sheetnames:
            return _failure("Add a sheet named Holidays, or use the downloadable template.")
        sheet = workbook["Holidays"]
        if sheet.max_column is not None and sheet.max_column > 128:
            return _failure("Use only the Date and Occasion columns in the Holidays sheet.")
        columns = max(2, sheet.max_column or 2)
        source_rows = sheet.iter_rows(max_row=MAX_IMPORT_ROWS + 2, max_col=columns)
        header = next(source_rows)
        headings = [str(cell.value or "").strip().lower() for cell in header[:2]]
        if headings != ["date", "occasion"]:
            return _failure("Row 1 must contain Date in column A and Occasion in column B.")
        if any(cell.value is not None for cell in header[2:]):
            return _failure("Use only the Date and Occasion columns in the Holidays sheet.")
        rows, errors, seen = [], [], {}
        for row_number, cells in enumerate(source_rows, start=2):
            values = [cell.value for cell in cells]
            if all(value is None or value == "" for value in values):
                continue
            if row_number > MAX_IMPORT_ROWS + 1:
                errors.append(_error("The workbook can contain at most 1,000 holiday rows.", row_number))
                break
            if any(cell.data_type == "f" for cell in cells):
                errors.append(_error("Replace formulas with their values before uploading.", row_number))
                continue
            if any(value is not None and value != "" for value in values[2:]):
                errors.append(_error("Use only the Date and Occasion columns.", row_number))
                continue
            try:
                day = _holiday_date(values[0])
            except (ValueError, OverflowError):
                errors.append(
                    _error("Enter a Gregorian Excel date or text in YYYY-MM-DD format.", row_number)
                )
                continue
            name = values[1].strip() if isinstance(values[1], str) else ""
            if not name:
                errors.append(_error("Enter an occasion name.", row_number))
                continue
            if len(name) > 160:
                errors.append(_error("The occasion name must be at most 160 characters.", row_number))
                continue
            if any(ord(character) < 32 for character in name):
                errors.append(_error("Use a single line of text for the occasion name.", row_number))
                continue
            if day in seen:
                errors.append(
                    _error(f"Duplicate date; this date also appears on row {seen[day]}.", row_number)
                )
                continue
            seen[day] = row_number
            rows.append({"row_number": row_number, "date": day.isoformat(), "name": name})
        if sheet.max_row is not None and sheet.max_row > MAX_IMPORT_ROWS + 2:
            errors.append(_error("The workbook can contain at most 1,000 holiday rows."))
        if not rows and not errors:
            errors.append(_error("Add at least one holiday below the header row."))
        return {"rows": rows, "errors": errors}
    except HolidayImportError as exc:
        return _failure(str(exc))
    except (BadZipFile, OSError, ValueError, KeyError, TypeError, IndexError, AttributeError):
        return _failure("This is not a readable Excel workbook. Save a new .xlsx file using the template.")
    except Exception:
        # The file format boundary can raise XML/backend-specific parsing errors.
        return _failure("This is not a readable Excel workbook. Save a new .xlsx file using the template.")
    finally:
        if workbook is not None:
            workbook.close()


def _validated_rows(rows):
    if not isinstance(rows, list) or not rows or len(rows) > MAX_IMPORT_ROWS:
        raise HolidayImportError("Upload between 1 and 1,000 holiday rows.")
    clean, seen = [], set()
    for row in rows:
        if not isinstance(row, dict):
            raise HolidayImportError("Upload and preview the workbook again.")
        try:
            day = _holiday_date(row.get("date"))
            name = row.get("name")
            number = row.get("row_number")
            valid = (
                isinstance(name, str)
                and bool(name.strip())
                and len(name.strip()) <= 160
                and not any(ord(character) < 32 for character in name)
                and type(number) is int
                and 2 <= number <= MAX_IMPORT_ROWS + 1
                and day not in seen
            )
        except (ValueError, TypeError, OverflowError):
            valid = False
        if not valid:
            raise HolidayImportError("Upload and preview the workbook again.")
        seen.add(day)
        clean.append({"row_number": number, "date": day.isoformat(), "name": name.strip()})
    return clean


def _day_start(day):
    return datetime.combine(day, time.min, LOCAL_TZ)


def _day_end(day):
    if day == date.max:
        return datetime.combine(day, time.max, LOCAL_TZ)
    return _day_start(day + timedelta(days=1))


def build_holiday_preview(rows):
    """Describe new/updated dates and affected active meetings without changing them."""
    rows = _validated_rows(rows)
    days = sorted(date.fromisoformat(row["date"]) for row in rows)
    holidays = {item.date: item for item in CompanyHoliday.objects.filter(date__in=days).order_by("date")}
    counts = {day: 0 for day in days}
    affected = []
    meetings = (
        Reservation.objects.filter(
            kind=Reservation.Kind.BOOKING,
            status__in=REVIEW_STATUSES,
            starts_at__lt=_day_end(days[-1]),
            ends_at__gt=_day_start(days[0]),
        )
        .order_by("pk")
        .values("pk", "starts_at", "ends_at", "status", "revision", "updated_at")
    )
    for meeting in meetings.iterator(chunk_size=1000):
        starts = meeting["starts_at"].astimezone(LOCAL_TZ)
        ends = meeting["ends_at"].astimezone(LOCAL_TZ)
        # A meeting ending at midnight belongs only to the preceding day.
        last_day = (ends - timedelta(microseconds=1)).date()
        touched = days[bisect_left(days, starts.date()) : bisect_right(days, last_day)]
        if not touched:
            continue
        for day in touched:
            counts[day] += 1
        affected.append(
            [
                meeting["pk"],
                starts.isoformat(),
                ends.isoformat(),
                meeting["status"],
                meeting["revision"],
                meeting["updated_at"].isoformat(),
            ]
        )
    summary = {"new": 0, "update": 0, "unchanged": 0, "affected_bookings": len(affected)}
    preview_rows, states = [], []
    for row in rows:
        day = date.fromisoformat(row["date"])
        holiday = holidays.get(day)
        action = "new" if holiday is None else "unchanged" if holiday.name == row["name"] else "update"
        summary[action] += 1
        preview_rows.append(
            {
                **row,
                "action": action,
                "previous_name": holiday.name if holiday else None,
                "booking_count": counts[day],
            }
        )
        states.append([row, holiday.pk if holiday else None, holiday.name if holiday else None])
    snapshot_data = json.dumps([states, affected], sort_keys=True, separators=(",", ":"))
    snapshot = hashlib.sha256(snapshot_data.encode("utf-8")).hexdigest()
    return {"rows": preview_rows, "summary": summary, "snapshot": snapshot}


def apply_holiday_import(rows, actor, expected_snapshot):
    """Apply a reviewed batch atomically; preserve other holidays and every booking."""
    rows = _validated_rows(rows)
    if not actor.is_authenticated or not actor.is_active or not actor.is_staff:
        raise HolidayImportError("Staff access is required.")
    try:
        with transaction.atomic():
            # Booking creates and moves lock rooms before checking holidays.
            # Hold the same locks through the final preview and import so an
            # ordinary booking cannot slip onto a newly closed date between them.
            # PK order matches multi-room booking updates; actor locks are held
            # by the staff view before this service is called.
            list(Room.objects.select_for_update().order_by("pk").values_list("pk", flat=True))
            dates = [date.fromisoformat(row["date"]) for row in rows]
            locked = {
                holiday.date: holiday
                for holiday in CompanyHoliday.objects.select_for_update()
                .filter(date__in=dates)
                .order_by("date")
            }
            preview = build_holiday_preview(rows)
            if not expected_snapshot or preview["snapshot"] != expected_snapshot:
                raise HolidayImportChanged(
                    "Holiday or booking data changed. Review a fresh preview before importing."
                )
            for row in preview["rows"]:
                if row["action"] == "unchanged":
                    continue
                day = date.fromisoformat(row["date"])
                holiday = locked.get(day)
                if holiday is None:
                    holiday = CompanyHoliday.objects.create(date=day, name=row["name"])
                else:
                    holiday.name = row["name"]
                    holiday.save(update_fields=["name"])
                AuditEvent.objects.create(
                    actor=actor,
                    action="holiday_saved",
                    target_type="holiday",
                    target_id=holiday.pk,
                    outcome="success",
                    details={
                        "source": "excel_import",
                        "date": row["date"],
                        "name": row["name"],
                        "previous_name": row["previous_name"],
                        "change": row["action"],
                    },
                )
            AuditEvent.objects.create(
                actor=actor,
                action="holiday_imported",
                target_type="holiday_import",
                outcome="success",
                details={**preview["summary"], "total": len(rows)},
            )
            return preview["summary"]
    except IntegrityError as exc:
        raise HolidayImportChanged("Holiday data changed. Review a fresh preview before importing.") from exc
